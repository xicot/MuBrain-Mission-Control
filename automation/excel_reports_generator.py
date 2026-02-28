#!/usr/bin/env python3
"""
Excel Reports Generator
Creates Excel files with clinical data and evolution charts
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference, BarChart
from datetime import datetime, timedelta
from pathlib import Path
import json
import sys
import requests

sys.path.insert(0, str(Path(__file__).parent))

# Config
AIRTABLE_TOKEN = ""
AIRTABLE_BASE_ID = "apphR8DK6lVJeY8n0"
SESSIONS_TABLE_ID = "tblFywptl9YawXrQJ"
PATIENTS_TABLE_ID = "tbldTQ3PxG5pcZVZ1"

# Load token from config
config_path = Path(__file__).parent / 'config.json'
if config_path.exists():
    with open(config_path, 'r') as f:
        config = json.load(f)
        AIRTABLE_TOKEN = config.get('airtable_token', '')

PRICING = {
    'Dr. FMT': {'NFB Online': {'invoice': 75, 'pay': 75}},
    'Dr. VA': {
        'NFB Presencial': {'invoice': 75, 'pay': 28},
        'QEEG': {'invoice': 180, 'pay': 15},
        'Somatic Session': {'invoice': 85, 'pay': 30},
        'Clinical Session': {'invoice': 65, 'pay': 35}
    }
}

def get_sessions_from_airtable(start_date, end_date):
    """Get completed sessions from Airtable"""
    url = f"https://api.airtable.com/v0/{AIRTABLE_BASE_ID}/{SESSIONS_TABLE_ID}"
    headers = {"Authorization": f"Bearer {AIRTABLE_TOKEN}"}
    params = {
        'filterByFormula': f"AND({{Session Completed}}, IS_AFTER({{Session Date & Time}}, '{start_date}'), IS_BEFORE({{Session Date & Time}}, '{end_date}T23:59:59'))"
    }
    all_records = []
    offset = None
    while True:
        if offset:
            params['offset'] = offset
        response = requests.get(url, headers=headers, params=params)
        if response.status_code == 200:
            data = response.json()
            all_records.extend(data.get('records', []))
            offset = data.get('offset')
            if not offset:
                break
        else:
            break
    return all_records

MEMORY_DIR = Path("C:/Users/MuLabs/.openclaw/workspace/memory")
REPORTS_DIR = MEMORY_DIR / "clinical-reports"
REPORTS_DIR.mkdir(exist_ok=True)


def get_weekly_metrics(weeks=12):
    """Get metrics for the last N weeks"""
    today = datetime.now().date()
    metrics = []
    
    for i in range(weeks):
        end_date = today - timedelta(weeks=i)
        start_date = end_date - timedelta(days=6)
        
        sessions = get_sessions_from_airtable(start_date.isoformat(), end_date.isoformat())
        
        total_sessions = len(sessions)
        total_invoice = 0
        total_pay = 0
        sessions_by_doctor = {'Dr. FMT': 0, 'Dr. VA': 0}
        
        for record in sessions:
            fields = record.get('fields', {})
            doctor = fields.get('Doctor Attended', 'Unknown')
            session_type = fields.get('Session Type', 'Unknown')
            
            if doctor in PRICING and session_type in PRICING[doctor]:
                total_invoice += PRICING[doctor][session_type]['invoice']
                total_pay += PRICING[doctor][session_type]['pay']
            
            if doctor in sessions_by_doctor:
                sessions_by_doctor[doctor] += 1
        
        metrics.append({
            'week_start': start_date,
            'week_end': end_date,
            'total_sessions': total_sessions,
            'total_invoice': total_invoice,
            'total_pay': total_pay,
            'net_margin': total_invoice - total_pay,
            'xico_sessions': sessions_by_doctor['Dr. FMT'],
            'vanessa_sessions': sessions_by_doctor['Dr. VA']
        })
    
    return list(reversed(metrics))  # Oldest first


def get_monthly_metrics(months=6):
    """Get metrics for the last N months"""
    today = datetime.now().date()
    metrics = []
    
    for i in range(months):
        # Calculate month
        year = today.year
        month = today.month - i
        while month <= 0:
            month += 12
            year -= 1
        
        # Get first and last day of month
        first_day = datetime(year, month, 1).date()
        if month == 12:
            last_day = datetime(year + 1, 1, 1).date() - timedelta(days=1)
        else:
            last_day = datetime(year, month + 1, 1).date() - timedelta(days=1)
        
        sessions = get_sessions_from_airtable(first_day.isoformat(), last_day.isoformat())
        
        total_sessions = len(sessions)
        total_invoice = 0
        total_pay = 0
        
        for record in sessions:
            fields = record.get('fields', {})
            doctor = fields.get('Doctor Attended', 'Unknown')
            session_type = fields.get('Session Type', 'Unknown')
            
            if doctor in PRICING and session_type in PRICING[doctor]:
                total_invoice += PRICING[doctor][session_type]['invoice']
                total_pay += PRICING[doctor][session_type]['pay']
        
        metrics.append({
            'month': f"{year}-{month:02d}",
            'month_name': first_day.strftime('%B %Y'),
            'total_sessions': total_sessions,
            'total_invoice': total_invoice,
            'total_pay': total_pay,
            'net_margin': total_invoice - total_pay
        })
    
    return list(reversed(metrics))


def create_excel_report():
    """Create comprehensive Excel report with charts"""
    
    # Create workbook
    wb = Workbook()
    
    # Remove default sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    
    # Weekly Dashboard
    ws_weekly = wb.create_sheet("Weekly Data")
    weekly_data = get_weekly_metrics(12)
    
    # Headers
    headers = ['Week Start', 'Week End', 'Sessions', 'Invoice', 'Pay', 'Net Margin', 'Xico', 'Vanessa']
    ws_weekly.append(headers)
    
    # Style headers
    for cell in ws_weekly[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0099CC", end_color="0099CC", fill_type="solid")
        cell.alignment = Alignment(horizontal='center')
    
    # Add weekly data
    for m in weekly_data:
        ws_weekly.append([
            m['week_start'].strftime('%Y-%m-%d'),
            m['week_end'].strftime('%Y-%m-%d'),
            m['total_sessions'],
            m['total_invoice'],
            m['total_pay'],
            m['net_margin'],
            m['xico_sessions'],
            m['vanessa_sessions']
        ])
    
    # Add charts
    # Sessions trend
    chart1 = LineChart()
    chart1.title = "Sessions per Week"
    chart1.style = 13
    chart1.y_axis.title = 'Sessions'
    chart1.x_axis.title = 'Week'
    
    data = Reference(ws_weekly, min_col=3, min_row=1, max_row=len(weekly_data)+1, max_col=3)
    cats = Reference(ws_weekly, min_col=1, min_row=2, max_row=len(weekly_data)+1)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    ws_weekly.add_chart(chart1, "J2")
    
    # Revenue trend
    chart2 = LineChart()
    chart2.title = "Revenue Evolution (€)"
    chart2.style = 13
    chart2.y_axis.title = 'Amount (€)'
    chart2.x_axis.title = 'Week'
    
    data = Reference(ws_weekly, min_col=4, min_row=1, max_row=len(weekly_data)+1, max_col=6)
    chart2.add_data(data, titles_from_data=True)
    chart2.set_categories(cats)
    ws_weekly.add_chart(chart2, "J18")
    
    # Adjust column widths
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']:
        ws_weekly.column_dimensions[col].width = 15
    
    # Monthly Dashboard
    ws_monthly = wb.create_sheet("Monthly Data")
    monthly_data = get_monthly_metrics(6)
    
    # Headers
    headers = ['Month', 'Month Name', 'Sessions', 'Invoice', 'Pay', 'Net Margin']
    ws_monthly.append(headers)
    
    for cell in ws_monthly[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="00D4FF", end_color="00D4FF", fill_type="solid")
        cell.alignment = Alignment(horizontal='center')
    
    # Add monthly data
    for m in monthly_data:
        ws_monthly.append([
            m['month'],
            m['month_name'],
            m['total_sessions'],
            m['total_invoice'],
            m['total_pay'],
            m['net_margin']
        ])
    
    # Monthly bar chart
    chart3 = BarChart()
    chart3.type = "col"
    chart3.style = 10
    chart3.title = "Monthly Revenue Breakdown"
    chart3.y_axis.title = 'Amount (€)'
    chart3.x_axis.title = 'Month'
    
    data = Reference(ws_monthly, min_col=4, min_row=1, max_row=len(monthly_data)+1, max_col=6)
    cats = Reference(ws_monthly, min_col=2, min_row=2, max_row=len(monthly_data)+1)
    chart3.add_data(data, titles_from_data=True)
    chart3.set_categories(cats)
    ws_monthly.add_chart(chart3, "H2")
    
    for col in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws_monthly.column_dimensions[col].width = 18
    
    # Summary Sheet
    ws_summary = wb.create_sheet("Summary", 0)
    ws_summary.append(["MuHealth Clinical Summary"])
    ws_summary.append([])
    
    # Calculate totals
    total_revenue = sum(m['total_invoice'] for m in monthly_data)
    total_sessions = sum(m['total_sessions'] for m in monthly_data)
    avg_weekly = weekly_data[-1]['total_sessions'] if weekly_data else 0
    
    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["Total Revenue (Last 6 Months)", f"€{total_revenue}"])
    ws_summary.append(["Total Sessions (Last 6 Months)", total_sessions])
    ws_summary.append(["Average Sessions/Week", f"{sum(m['total_sessions'] for m in weekly_data) / len(weekly_data):.1f}"])
    ws_summary.append(["Current Week Sessions", weekly_data[-1]['total_sessions'] if weekly_data else 0])
    ws_summary.append(["Net Margin (Last 6 Months)", f"€{sum(m['net_margin'] for m in monthly_data)}"])
    
    # Style summary
    ws_summary['A1'].font = Font(size=16, bold=True, color="0099CC")
    ws_summary['A3'].font = Font(bold=True)
    ws_summary['B3'].font = Font(bold=True)
    
    ws_summary.column_dimensions['A'].width = 35
    ws_summary.column_dimensions['B'].width = 25
    
    # Save file
    filename = f"clinical-evolution-{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    filepath = REPORTS_DIR / filename
    wb.save(filepath)
    
    print(f"[OK] Excel report created: {filepath}")
    return filepath, {
        'total_revenue': total_revenue,
        'total_sessions': total_sessions,
        'weekly_data': weekly_data,
        'monthly_data': monthly_data
    }


def upload_excel_to_drive(filepath):
    """Upload Excel report to Google Drive"""
    try:
        sys.path.insert(0, str(Path(__file__).parent))
        from drive_upload_helper import upload_report
        
        file_id, link = upload_report(
            str(filepath),
            filepath.name,
            'Weekly and monthly clinical evolution with charts'
        )
        return file_id, link
    except Exception as e:
        print(f"[ERROR] Upload failed: {e}")
        return None, None


def generate_and_upload_excel_report():
    """Generate Excel report and upload to Drive"""
    print("="*60)
    print("GENERATING EXCEL EVOLUTION REPORT")
    print("="*60)
    
    # Generate report
    result = create_excel_report()
    
    if not result:
        print("[ERROR] Failed to generate report")
        return None
    
    filepath, metadata = result
    
    # Upload to Drive
    print("\n[UPLOAD] Sending to Google Drive...")
    file_id, link = upload_excel_to_drive(filepath)
    
    if link:
        print(f"[OK] Excel uploaded: {link}")
        
        # Update manifest
        manifest_file = MEMORY_DIR / "reports-manifest.json"
        manifest = {}
        if manifest_file.exists():
            with open(manifest_file, 'r') as f:
                manifest = json.load(f)
        
        if 'excel_reports' not in manifest:
            manifest['excel_reports'] = []
        
        manifest['excel_reports'].insert(0, {
            'date': datetime.now().isoformat(),
            'filename': filepath.name,
            'local_path': str(filepath),
            'drive_link': link,
            'metadata': {
                'total_revenue': metadata['total_revenue'],
                'total_sessions': metadata['total_sessions']
            }
        })
        manifest['excel_reports'] = manifest['excel_reports'][:10]  # Keep last 10
        
        with open(manifest_file, 'w') as f:
            json.dump(manifest, f, indent=2)
        
        return {
            'filepath': str(filepath),
            'drive_link': link,
            'metadata': metadata
        }
    
    return None


if __name__ == "__main__":
    result = generate_and_upload_excel_report()
    if result:
        print(f"\n[COMPLETE] Excel report ready!")
        print(f"Drive: {result['drive_link']}")
